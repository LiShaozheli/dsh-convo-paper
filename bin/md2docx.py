#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Minimal Markdown -> DOCX converter tuned for academic Chinese papers.

Supported: #/##/### headings, paragraphs, **bold**, *italic*, `code`,
blockquotes, ordered/unordered lists, simple tables (pipe style), horizontal
rules. Chinese body text uses 宋体 小四 / 1.5 spacing / 2-char first-line indent;
headings use 黑体.

Usage:
    python md2docx.py paper.md paper.docx [--title "自定义标题"]
"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt

BODY_FONT = "宋体"
HEAD_FONT = "黑体"
LATIN_FONT = "Times New Roman"


def set_run_font(run, name=BODY_FONT, size=None, bold=None):
    run.font.name = LATIN_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


def add_inline(paragraph, text, size=12):
    """Render **bold**, *italic* and `code` inside one paragraph."""
    pattern = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")
    for chunk in pattern.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            run = paragraph.add_run(chunk[2:-2])
            set_run_font(run, size=size, bold=True)
        elif chunk.startswith("`") and chunk.endswith("`"):
            run = paragraph.add_run(chunk[1:-1])
            set_run_font(run, name="Consolas", size=size - 0.5)
        elif chunk.startswith("*") and chunk.endswith("*"):
            run = paragraph.add_run(chunk[1:-1])
            set_run_font(run, size=size)
            run.italic = True
        else:
            run = paragraph.add_run(chunk)
            set_run_font(run, size=size)


def body_paragraph(doc, text, indent=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    if indent:
        pf.first_line_indent = Pt(24)  # 2 chars at 12pt
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(p, text)
    return p


def heading(doc, text, level):
    size = {1: 18, 2: 14, 3: 12.5}.get(level, 12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12 if level > 1 else 0)
    p.paragraph_format.space_after = Pt(8 if level > 1 else 18)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = p.add_run(text)
    set_run_font(run, name=HEAD_FONT, size=size, bold=True)
    return p


def parse_table_row(line):
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells


def convert(src, dst, title=None):
    lines = [l.rstrip("\n") for l in open(src, encoding="utf-8").read().split("\n")]
    doc = Document()

    # Page setup: A4 with normal margins.
    section = doc.sections[0]
    section.page_width = Pt(595)
    section.page_height = Pt(842)

    normal = doc.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    i = 0
    first_h1 = True
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("```"):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                i += 1
            i += 1
            continue

        m = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1 and title:
                text = title
            heading(doc, text, level)
            first_h1 = False
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            header = parse_table_row(stripped)
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(parse_table_row(lines[i]))
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            for c, text in enumerate(header):
                cell = table.rows[0].cells[c]
                cell.text = ""
                add_inline(cell.paragraphs[0], text, size=10.5)
            for row in rows:
                cells = table.add_row().cells
                for c, text in enumerate(row[: len(header)]):
                    cells[c].text = ""
                    add_inline(cells[c].paragraphs[0], text, size=10.5)
            doc.add_paragraph()
            continue

        if stripped.startswith(">"):
            p = body_paragraph(doc, stripped.lstrip("> ").strip(), indent=False)
            p.paragraph_format.left_indent = Pt(18)
            for run in p.runs:
                run.font.color.rgb = None
            i += 1
            continue

        m = re.match(r"^(\d+)[.)]\s+(.*)$", stripped)
        if m:
            p = body_paragraph(doc, m.group(2), indent=False)
            p.paragraph_format.left_indent = Pt(24)
            p.paragraph_format.first_line_indent = Pt(-12)
            i += 1
            continue

        if stripped.startswith(("- ", "* ")):
            p = body_paragraph(doc, stripped[2:], indent=False)
            p.paragraph_format.left_indent = Pt(24)
            p.paragraph_format.first_line_indent = Pt(-12)
            i += 1
            continue

        if re.match(r"^(-{3,}|\*{3,})$", stripped):
            doc.add_paragraph()
            i += 1
            continue

        # Merge hard-wrapped lines into one paragraph.
        buf = [stripped]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(
            r"^(#{1,3}\s|\||>|\d+[.)]\s|[-*]\s|```)", lines[i].strip()
        ):
            buf.append(lines[i].strip())
            i += 1
        body_paragraph(doc, " ".join(buf) if not buf[0].startswith("[") else " ".join(buf))

    doc.save(dst)
    return dst


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    title = None
    if "--title" in sys.argv:
        title = sys.argv[sys.argv.index("--title") + 1]
    print(convert(sys.argv[1], sys.argv[2], title))
